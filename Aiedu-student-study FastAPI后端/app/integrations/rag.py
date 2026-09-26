from __future__ import annotations

import hashlib
import io
import logging
import math
import re
from collections import Counter
from collections.abc import Iterable, Sequence
from dataclasses import asdict, dataclass
from functools import lru_cache
from time import perf_counter
from typing import Any, Literal

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings


logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class DocumentBlock:
    block_type: str
    text: str
    heading_path: str | None = None
    page_number: int | None = None
    slide_number: int | None = None


RetrievalMode = Literal["bm25", "dense", "hybrid", "hybrid_rerank"]


@dataclass(frozen=True)
class RetrievalResult:
    """Inspectable retrieval output shared by production and offline evaluation."""

    mode: RetrievalMode
    results: list[dict[str, Any]]
    timings_ms: dict[str, float]
    diagnostics: dict[str, Any]


_CHINESE_NUMBERS = {
    "一": 1, "二": 2, "三": 3, "四": 4, "五": 5,
    "六": 6, "七": 7, "八": 8, "九": 9, "十": 10,
}
_CHAPTER_PATTERN = re.compile(
    r"^第\s*(?P<number>[一二三四五六七八九十百\d]+)\s*章\s*(?P<title>[^【]{0,80})"
)
_SECTION_PATTERN = re.compile(
    r"^(?P<number>\d+(?:\s*\.\s*\d+)+)\s*\.?\s*(?P<title>\S.{0,80})$"
)
_CHINESE_HEADING_PATTERN = re.compile(r"^(?P<number>[一二三四五六七八九十]+)、\s*(?P<title>\S.{0,80})$")
_BRACKET_HEADING_PATTERN = re.compile(r"^[【\[](?P<title>[^】\]]{2,40})[】\]]$")
_PDF_NOISE_PATTERN = re.compile(
    r"(?:www\.[a-z0-9.-]+\.[a-z]{2,}|Linux公社\s*\(LinuxIDC\.com\).*)",
    re.IGNORECASE,
)


def _heading_info(value: str) -> tuple[int, str] | None:
    """Recognize the heading conventions used by the uploaded Chinese course files."""
    line = re.sub(r"\s+", " ", value).strip()
    if not line or len(line) > 120 or re.search(r"\.{4,}\s*\d*\s*$", line):
        return None
    match = _CHAPTER_PATTERN.match(line)
    if match:
        title = match.group("title").strip(" ：:。")
        if (len(title) > 36 or re.search(r"[。！？；]", line) or
                re.search(r"第\s*[一二三四五六七八九十百\d]+\s*章", title) or
                re.match(r"^(?:的|将|中|里|对此)", title)):
            return None
        return 1, f"第{match.group('number')}章" + (f" {title}" if title else "")
    match = _SECTION_PATTERN.match(line)
    if match:
        number = re.sub(r"\s+", "", match.group("number"))
        title = match.group("title").strip()
        first_number = int(number.split(".", 1)[0])
        if first_number > 20 or re.search(r"[。！？；]", title):
            return None
        return 2, f"{number} {title}"
    match = _CHINESE_HEADING_PATTERN.match(line)
    if match:
        return 1, f"{match.group('number')}、{match.group('title').strip()}"
    match = _BRACKET_HEADING_PATTERN.match(line)
    if match:
        return 3, match.group("title").strip()
    return None


def _update_heading_path(headings: list[str], level: int, value: str) -> str:
    if level <= 1:
        headings[:] = [value]
    else:
        headings[:] = headings[:level - 1]
        while len(headings) < level - 1:
            headings.append("")
        headings.append(value)
    return " > ".join(item for item in headings if item)


def _clean_pdf_line(value: str) -> str:
    line = re.sub(r"[ \t]+", " ", value).strip()
    if not line or _PDF_NOISE_PATTERN.search(line):
        return ""
    if re.fullmatch(r"[-–—]?\s*\d{1,4}\s*[-–—]?", line):
        return ""
    return line


def _pdf_repeated_edge_lines(pages: Sequence[str]) -> set[str]:
    counts: Counter[str] = Counter()
    for text in pages:
        lines = [_clean_pdf_line(line) for line in text.splitlines()]
        lines = [line for line in lines if line]
        for line in dict.fromkeys(lines[:2] + lines[-3:]):
            if (_heading_info(line) is None and len(line) >= 6 and
                    not re.search(r"[{};=#<>]", line)):
                counts[line] += 1
    threshold = max(3, math.ceil(len(pages) * 0.08))
    return {line for line, count in counts.items() if count >= threshold}


def _is_toc_page(lines: Sequence[str]) -> bool:
    dotted = sum(bool(re.search(r"\.{4,}\s*\d+\s*$", line)) for line in lines)
    return dotted >= 3


def _section_index_blocks(blocks: Sequence[DocumentBlock]) -> list[DocumentBlock]:
    chapters: dict[str, dict[str, Any]] = {}
    for block in blocks:
        if block.block_type != "heading" or not block.heading_path:
            continue
        parts = [item.strip() for item in block.heading_path.split(">") if item.strip()]
        if len(parts) == 1 and _CHAPTER_PATTERN.match(parts[0]):
            chapters.setdefault(parts[0], {"page": block.page_number, "sections": []})
        elif len(parts) >= 2 and _CHAPTER_PATTERN.match(parts[0]):
            value = parts[1]
            if _SECTION_PATTERN.match(value) is None:
                continue
            chapter = chapters.setdefault(
                parts[0], {"page": block.page_number, "sections": []},
            )
            if value not in chapter["sections"]:
                chapter["sections"].append(value)
    values = []
    for chapter, data in chapters.items():
        sections = data["sections"]
        if sections:
            values.append(DocumentBlock(
                "section_index",
                f"{chapter}包含以下小节：" + "；".join(sections) + "。",
                chapter,
                page_number=data["page"],
            ))
    return values


def extract_blocks(content: bytes, mime_type: str) -> tuple[list[DocumentBlock], int | None]:
    """Extract format-aware blocks while retaining citation locators."""
    blocks: list[DocumentBlock] = []
    if mime_type == "application/pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        raw_pages = [page.extract_text() or "" for page in reader.pages]
        repeated_edges = _pdf_repeated_edge_lines(raw_pages)
        headings: list[str] = []
        for page_number, text in enumerate(raw_pages, 1):
            raw_lines = [line.strip() for line in text.splitlines() if line.strip()]
            if _is_toc_page(raw_lines):
                continue
            lines = [_clean_pdf_line(line) for line in raw_lines]
            lines = [line for line in lines if line and line not in repeated_edges]
            paragraph: list[str] = []

            def flush_paragraph() -> None:
                if not paragraph:
                    return
                value = " ".join(paragraph).strip()
                if value:
                    blocks.append(DocumentBlock(
                        "paragraph", value, " > ".join(item for item in headings if item) or None,
                        page_number=page_number,
                    ))
                paragraph.clear()

            for line in lines:
                heading = _heading_info(line)
                if heading:
                    flush_paragraph()
                    level, title = heading
                    path = _update_heading_path(headings, level, title)
                    blocks.append(DocumentBlock("heading", title, path, page_number=page_number))
                else:
                    paragraph.append(line)
            flush_paragraph()
        blocks.extend(_section_index_blocks(blocks))
        return blocks, len(reader.pages)
    if mime_type.endswith("wordprocessingml.document"):
        from docx import Document

        document = Document(io.BytesIO(content))
        headings: list[str] = []
        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            if not value:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style else ""
            detected = _heading_info(value)
            if style.startswith("heading") or style.startswith("标题"):
                match = re.search(r"(\d+)$", style)
                level = int(match.group(1)) if match else 1
                path = _update_heading_path(headings, level, value)
                blocks.append(DocumentBlock("heading", value, path))
            elif detected:
                level, title = detected
                path = _update_heading_path(headings, level, title)
                blocks.append(DocumentBlock("heading", title, path))
            else:
                block_type = "list" if "list" in style or "列表" in style else "paragraph"
                blocks.append(DocumentBlock(
                    block_type, value,
                    " > ".join(item for item in headings if item) or None,
                ))
        for table in document.tables:
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            value = "\n".join(row for row in rows if row.strip(" |"))
            if value:
                blocks.append(DocumentBlock(
                    "table", value,
                    " > ".join(item for item in headings if item) or None,
                ))
        blocks.extend(_section_index_blocks(blocks))
        return blocks, None
    if mime_type.endswith("presentationml.presentation"):
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(content))
        for slide_number, slide in enumerate(presentation.slides, 1):
            title = ""
            if slide.shapes.title is not None:
                title = (slide.shapes.title.text or "").strip()
            values: list[str] = []
            for shape in slide.shapes:
                if shape is slide.shapes.title:
                    continue
                if getattr(shape, "has_table", False):
                    rows = [" | ".join(cell.text.strip() for cell in row.cells)
                            for row in shape.table.rows]
                    table_text = "\n".join(row for row in rows if row.strip(" |"))
                    if table_text:
                        values.append(table_text)
                elif hasattr(shape, "text") and shape.text.strip():
                    values.append(shape.text.strip())
            text = "\n".join(values).strip()
            if title:
                blocks.append(DocumentBlock("heading", title, title, slide_number=slide_number))
            if text:
                blocks.append(DocumentBlock("slide", text, title or None, slide_number=slide_number))
        return blocks, len(presentation.slides)
    raise ValueError("不支持的课程资料类型")


def extract_text(content: bytes, mime_type: str) -> tuple[str, int | None]:
    blocks, page_count = extract_blocks(content, mime_type)
    return "\n\n".join(block.text for block in blocks), page_count


def _estimated_tokens(text: str) -> int:
    # Chinese is close to one token per character; latin text is closer to 4 chars/token.
    chinese = len(re.findall(r"[\u4e00-\u9fff]", text))
    return chinese + max(1, (len(text) - chinese) // 4)


def _prefix_for_tokens(text: str, token_limit: int) -> int:
    low, high = 1, len(text)
    while low < high:
        middle = (low + high + 1) // 2
        if _estimated_tokens(text[:middle]) <= token_limit:
            low = middle
        else:
            high = middle - 1
    return low


def _suffix_for_tokens(text: str, token_limit: int) -> str:
    if token_limit <= 0:
        return ""
    if _estimated_tokens(text) <= token_limit:
        return text
    low, high = 1, len(text)
    while low < high:
        middle = (low + high + 1) // 2
        if _estimated_tokens(text[-middle:]) <= token_limit:
            low = middle
        else:
            high = middle - 1
    return text[-low:].lstrip()


def _token_bounded_parts(text: str, maximum: int, overlap: int) -> list[str]:
    remaining = text.strip()
    values: list[str] = []
    while remaining:
        if _estimated_tokens(remaining) <= maximum:
            values.append(remaining)
            break
        end = _prefix_for_tokens(remaining, maximum)
        lower = max(1, end * 3 // 5)
        boundaries = [remaining.rfind(mark, lower, end) for mark in ("\n", "。", "！", "？", ". ", "; ")]
        natural = max(boundaries)
        if natural >= lower:
            end = natural + 1
        value = remaining[:end].strip()
        values.append(value)
        tail = _suffix_for_tokens(value, overlap)
        remaining = (tail + remaining[end:]).strip()
        if remaining == value:
            remaining = remaining[end:].strip()
    return values


def chunk_blocks(blocks: Sequence[DocumentBlock]) -> list[dict[str, Any]]:
    """Create section-aware chunks with bounded overlap and stable locator metadata."""
    settings = get_settings()
    target = settings.rag_chunk_target_tokens
    maximum = settings.rag_chunk_max_tokens
    overlap = settings.rag_chunk_overlap_tokens
    output: list[dict[str, Any]] = []
    current: list[str] = []
    current_tokens = 0
    meta: DocumentBlock | None = None

    def flush() -> None:
        nonlocal current, current_tokens, meta
        if not current or meta is None:
            return
        text = "\n\n".join(current).strip()
        if text:
            output.append({**asdict(meta), "text": text, "token_count": _estimated_tokens(text)})
        tail = _suffix_for_tokens(text, overlap)
        current = [tail] if tail else []
        current_tokens = _estimated_tokens(tail) if tail else 0

    for block in blocks:
        paragraphs = [value.strip() for value in re.split(r"\n\s*\n+", block.text) if value.strip()]
        if not paragraphs:
            continue
        if meta is not None and (block.heading_path != meta.heading_path or
                                 block.page_number != meta.page_number or
                                 block.slide_number != meta.slide_number):
            same_pdf_section = (
                block.heading_path == meta.heading_path and
                block.page_number is not None and meta.page_number is not None and
                block.slide_number is None and meta.slide_number is None
            )
            flush()
            # A PDF page break is a layout boundary, not a semantic boundary.
            # Keep only the configured overlap when the same section continues
            # on the next page so citations do not begin with an orphaned clause.
            if not same_pdf_section:
                current = []
                current_tokens = 0
        meta = block
        for paragraph in paragraphs:
            parts = _token_bounded_parts(paragraph, max(1, maximum - overlap), 0)
            for part in parts:
                part_tokens = _estimated_tokens(part)
                if current and current_tokens + part_tokens > maximum:
                    flush()
                    meta = block
                current.append(part)
                current_tokens += part_tokens
                if current_tokens >= target:
                    flush()
                    meta = block
    flush()
    # A parent heading immediately followed by a child heading has no evidence
    # of its own. Its text already survives in the child's heading_path, so do
    # not index tiny heading-only records that would crowd out answer passages.
    values = [item for item in output if item.get("block_type") != "heading"]
    merged: list[dict[str, Any]] = []
    for item in values:
        if (item["token_count"] < 30 and merged and
                item.get("heading_path") == merged[-1].get("heading_path") and
                item.get("page_number") == merged[-1].get("page_number") and
                item.get("slide_number") == merged[-1].get("slide_number") and
                merged[-1]["token_count"] + item["token_count"] <= maximum):
            merged[-1]["text"] = f'{merged[-1]["text"]}\n\n{item["text"]}'
            merged[-1]["token_count"] = _estimated_tokens(merged[-1]["text"])
        else:
            merged.append(item)
    return merged


def chunks(text: str, size: int | None = None, overlap: int | None = None) -> Iterable[str]:
    """Split at natural boundaries where possible and retain semantic overlap."""
    settings = get_settings()
    size = size or settings.rag_chunk_size_chars
    overlap = settings.rag_chunk_overlap_chars if overlap is None else overlap
    if size <= 0 or overlap < 0 or overlap >= size:
        raise ValueError("切片大小配置无效")

    normalized = re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n").replace("\r", "\n"))
    normalized = re.sub(r"\n{3,}", "\n\n", normalized).strip()
    start = 0
    while start < len(normalized):
        end = min(start + size, len(normalized))
        if end < len(normalized):
            lower_bound = start + max(size // 2, 1)
            candidates = [
                normalized.rfind("\n\n", lower_bound, end),
                normalized.rfind("。", lower_bound, end),
                normalized.rfind("！", lower_bound, end),
                normalized.rfind("？", lower_bound, end),
                normalized.rfind(". ", lower_bound, end),
            ]
            boundary = max(candidates)
            if boundary >= lower_bound:
                end = boundary + (1 if normalized[boundary] != "\n" else 2)
        value = normalized[start:end].strip()
        if value:
            yield value
        if end >= len(normalized):
            break
        start = max(start + 1, end - overlap)


def embed_texts(texts: Sequence[str]) -> list[list[float]]:
    """Call an OpenAI-compatible embeddings endpoint in bounded batches."""
    settings = get_settings()
    if not settings.ai_api_key:
        raise RuntimeError("尚未配置 AIEDU_AI_API_KEY，资料已保存，可配置后重新索引")
    if settings.embedding_model != "text-embedding-3-large":
        raise RuntimeError("课程知识库当前仅允许使用 text-embedding-3-large")

    from openai import OpenAI

    client = OpenAI(
        api_key=settings.ai_api_key,
        base_url=settings.ai_base_url,
        timeout=60,
        max_retries=3,
    )
    vectors: list[list[float]] = []
    batch_size = max(1, settings.embedding_batch_size)
    for offset in range(0, len(texts), batch_size):
        batch = list(texts[offset:offset + batch_size])
        batch_number = offset // batch_size + 1
        batch_total = math.ceil(len(texts) / batch_size)
        started = perf_counter()
        logger.info(
            "Embedding batch %s/%s: %s chunks with %s",
            batch_number, batch_total, len(batch), settings.embedding_model,
        )
        try:
            response = client.embeddings.create(
                model=settings.embedding_model,
                input=batch,
                dimensions=settings.embedding_dimensions,
            )
        except Exception:
            logger.exception("Embedding batch %s/%s failed", batch_number, batch_total)
            raise
        logger.info(
            "Embedding batch %s/%s completed in %sms",
            batch_number, batch_total, round((perf_counter() - started) * 1000),
        )
        ordered = sorted(response.data, key=lambda item: item.index)
        if len(ordered) != len(batch):
            raise RuntimeError("Embedding 服务返回的向量数量不匹配")
        for item in ordered:
            vector = list(item.embedding)
            if len(vector) != settings.embedding_dimensions:
                raise RuntimeError("Embedding 服务返回的向量维度不匹配")
            vectors.append(vector)
    return vectors


def _qdrant_client(timeout: int = 10):
    from qdrant_client import QdrantClient

    settings = get_settings()
    local = re.match(r"^https?://(?:localhost|127\.0\.0\.1|\[::1\])(?::|/|$)", settings.qdrant_url)
    return QdrantClient(
        url=settings.qdrant_url,
        timeout=timeout,
        # httpx otherwise consults the Windows proxy registry and may route a
        # localhost Qdrant request through the user's HTTP proxy.
        trust_env=not bool(local),
    )


def _ensure_collection(client) -> None:
    from qdrant_client import models

    settings = get_settings()
    collections = {item.name for item in client.get_collections().collections}
    if settings.qdrant_collection not in collections:
        client.create_collection(
            collection_name=settings.qdrant_collection,
            vectors_config=models.VectorParams(
                size=settings.embedding_dimensions,
                distance=models.Distance.COSINE,
            ),
        )
        return
    collection = client.get_collection(settings.qdrant_collection)
    actual_size = getattr(collection.config.params.vectors, "size", None)
    if actual_size != settings.embedding_dimensions:
        raise RuntimeError(
            f"Qdrant 集合向量维度为 {actual_size}，配置要求 {settings.embedding_dimensions}，请更换集合名后重试"
        )


def _chapter_aliases(value: str) -> str:
    def replace(match: re.Match[str]) -> str:
        number = match.group(1)
        if number.isdigit():
            reverse = {value: key for key, value in _CHINESE_NUMBERS.items()}
            chinese = reverse.get(int(number))
            return f"第{number}章" + (f" 第{chinese}章" if chinese else "")
        arabic = _CHINESE_NUMBERS.get(number)
        return f"第{number}章" + (f" 第{arabic}章" if arabic else "")

    return re.sub(r"第\s*([一二三四五六七八九十\d]+)\s*章", replace, value)


def _resource_type_label(resource_type: str | None) -> str:
    return {
        "syllabus": "课程大纲",
        "textbook": "教材",
        "courseware": "课件",
    }.get(resource_type or "", resource_type or "课程资料")


def _structured_retrieval_text(
    *, title: str, resource_type: str | None, heading_path: str | None, text: str,
) -> str:
    prefix = [f"资料类型：{_resource_type_label(resource_type)}", f"资料标题：{title}"]
    if heading_path:
        prefix.append(f"章节：{_chapter_aliases(heading_path)}")
    return "\n".join(prefix + [text])


def _candidate_retrieval_text(item: dict[str, Any]) -> str:
    return _structured_retrieval_text(
        title=str(item.get("title") or "课程资料"),
        resource_type=item.get("resource_type"),
        heading_path=item.get("heading_path"),
        text=str(item.get("text") or ""),
    )


def _requested_resource_types(query: str) -> list[str]:
    values: list[str] = []
    for marker, resource_type in (("教材", "textbook"), ("大纲", "syllabus"), ("课件", "courseware")):
        if marker in query:
            values.append(resource_type)
    # In this project, “课程第几章要求/教学目标” refers to the syllabus even
    # when the student does not literally say “大纲”. This is a domain routing
    # hint, not a generic semantic guess.
    if (re.search(r"课程.{0,24}(?:要求|目标|内容|学时)", query) or
            re.search(r"(?:教学|学习)(?:要求|目标)", query)):
        values.append("syllabus")
    return list(dict.fromkeys(values))


def _chapter_number(value: str) -> int | None:
    compact = re.sub(r"\s+", "", value)
    if compact.isdigit():
        return int(compact)
    if compact in _CHINESE_NUMBERS:
        return _CHINESE_NUMBERS[compact]
    if compact.startswith("十"):
        return 10 + _CHINESE_NUMBERS.get(compact[1:], 0)
    if compact.endswith("十"):
        return _CHINESE_NUMBERS.get(compact[:-1], 0) * 10
    if "十" in compact:
        tens, ones = compact.split("十", 1)
        return _CHINESE_NUMBERS.get(tens, 1) * 10 + _CHINESE_NUMBERS.get(ones, 0)
    return None


def _query_chapters(query: str) -> set[int]:
    values: set[int] = set()
    for match in re.finditer(r"第\s*([一二三四五六七八九十百\d]+)\s*章", query):
        number = _chapter_number(match.group(1))
        if number is not None:
            values.add(number)
    return values


def _candidate_chapter(item: dict[str, Any]) -> int | None:
    value = str(item.get("heading_path") or item.get("text") or "")[:160]
    match = re.search(r"第\s*([一二三四五六七八九十百\d]+)\s*章", value)
    return _chapter_number(match.group(1)) if match else None


def _apply_source_preferences(query: str, candidates: list[dict]) -> list[dict]:
    requested = _requested_resource_types(query)
    chapters = _query_chapters(query)
    matching = [item for item in candidates if item.get("resource_type") in requested]
    if len(requested) == 1:
        candidates = matching or candidates
    elif len(requested) > 1:
        candidates = matching or candidates

    def chapter_matches(item: dict) -> bool:
        return not chapters or _candidate_chapter(item) in chapters

    if not requested:
        return sorted(candidates, key=lambda item: not chapter_matches(item))

    if len(requested) > 1:
        queues: dict[str, list[dict]] = {}
        used: set[str] = set()
        for resource_type in requested:
            values = [item for item in candidates if item.get("resource_type") == resource_type]
            queues[resource_type] = sorted(values, key=lambda item: not chapter_matches(item))
        interleaved: list[dict] = []
        while any(queues.values()):
            for resource_type in requested:
                if not queues[resource_type]:
                    continue
                item = queues[resource_type].pop(0)
                key = str(item.get("chunk_id") or f"{item.get('resource_id')}:{item.get('position')}")
                if key not in used:
                    interleaved.append(item)
                    used.add(key)
        interleaved.extend(
            item for item in candidates
            if str(item.get("chunk_id") or f"{item.get('resource_id')}:{item.get('position')}") not in used
        )
        return interleaved

    promoted: list[dict] = []
    used: set[str] = set()
    for resource_type in requested:
        same_source = [row for row in candidates if row.get("resource_type") == resource_type]
        item = next((row for row in same_source if chapter_matches(row)), None)
        item = item or next(iter(same_source), None)
        if item is not None:
            key = str(item.get("chunk_id") or f"{item.get('resource_id')}:{item.get('position')}")
            promoted.append(item)
            used.add(key)
    remaining = [
        item for item in candidates
        if str(item.get("chunk_id") or f"{item.get('resource_id')}:{item.get('position')}") not in used
    ]
    promoted.extend(sorted(remaining, key=lambda item: not chapter_matches(item)))
    return promoted


def index_resource(
    resource_id: int, offering_id: int, title: str,
    text: str | None = None, indexed_chunks: Sequence[dict[str, Any]] | None = None,
    resource_type: str | None = None,
) -> int:
    from qdrant_client import models

    settings = get_settings()
    if indexed_chunks is None:
        values = [{"text": value, "position": position}
                  for position, value in enumerate(chunks(text or ""))]
    else:
        values = list(indexed_chunks)
    if not values:
        raise ValueError("资料未提取到可索引文本")

    # Do not delete a usable old index until every new vector exists.
    vectors = embed_texts([
        _structured_retrieval_text(
            title=title,
            resource_type=resource_type,
            heading_path=item.get("heading_path"),
            text=str(item["text"]),
        )
        for item in values
    ])
    client = _qdrant_client()
    _ensure_collection(client)
    points = []
    for position, (item, vector) in enumerate(zip(values, vectors, strict=True)):
        value = str(item["text"])
        point_id = int.from_bytes(
            hashlib.sha256(f"{resource_id}:{position}".encode()).digest()[:8], "big"
        ) >> 1
        points.append(models.PointStruct(
            id=point_id,
            vector=vector,
            payload={
                "resource_id": resource_id,
                "offering_id": offering_id,
                "title": title,
                "resource_type": resource_type,
                "position": position,
                "text": value,
                "chunk_id": item.get("id"),
                "block_type": item.get("block_type", "paragraph"),
                "heading_path": item.get("heading_path"),
                "page_number": item.get("page_number"),
                "slide_number": item.get("slide_number"),
                "embedding_model": settings.embedding_model,
            },
        ))
    resource_filter = models.Filter(must=[
        models.FieldCondition(key="resource_id", match=models.MatchValue(value=resource_id))
    ])
    old_point_ids: set[int | str] = set()
    offset = None
    while True:
        records, offset = client.scroll(
            collection_name=settings.qdrant_collection,
            scroll_filter=resource_filter,
            limit=256,
            offset=offset,
            with_payload=False,
            with_vectors=False,
        )
        old_point_ids.update(record.id for record in records)
        if offset is None:
            break

    # A 3072-dimensional embedding makes a few hundred points exceed Qdrant's
    # default 32 MiB HTTP body limit. Upsert in bounded batches and only remove
    # obsolete old point IDs after every new batch has succeeded. This also
    # prevents a transient upload failure from first erasing the usable index.
    batch_size = 64
    for start in range(0, len(points), batch_size):
        client.upsert(
            collection_name=settings.qdrant_collection,
            points=points[start:start + batch_size],
            wait=True,
        )
    new_point_ids = {point.id for point in points}
    stale_ids = list(old_point_ids - new_point_ids)
    for start in range(0, len(stale_ids), 256):
        client.delete(
            collection_name=settings.qdrant_collection,
            points_selector=models.PointIdsList(points=stale_ids[start:start + 256]),
            wait=True,
        )
    return len(points)


def _tokenize(text: str) -> list[str]:
    normalized = re.sub(r"\s+", "", text.lower())
    latin = re.findall(r"[a-z0-9_]+", normalized)
    chinese = re.findall(r"[\u4e00-\u9fff]", normalized)
    bigrams = ["".join(chinese[index:index + 2]) for index in range(max(0, len(chinese) - 1))]
    return latin + chinese + bigrams


def _bm25_search(db: Session, offering_id: int, query: str, limit: int) -> list[dict]:
    from app.models import CourseResource, ResourceChunk

    rows = db.execute(
        select(ResourceChunk, CourseResource.title, CourseResource.resource_type)
        .join(CourseResource, CourseResource.id == ResourceChunk.resource_id)
        .where(ResourceChunk.offering_id == offering_id, CourseResource.deleted_at.is_(None))
        .order_by(ResourceChunk.id)
    ).all()
    if not rows:
        return []
    documents = [_tokenize(_structured_retrieval_text(
        title=title, resource_type=resource_type,
        heading_path=chunk.heading_path, text=chunk.text,
    )) for chunk, title, resource_type in rows]
    query_tokens = _tokenize(_chapter_aliases(query))
    if not query_tokens:
        return []
    document_frequency: Counter[str] = Counter()
    for tokens in documents:
        document_frequency.update(set(tokens))
    average_length = sum(len(tokens) for tokens in documents) / max(1, len(documents))
    scored: list[tuple[float, Any, str, str | None]] = []
    for (chunk, title, resource_type), tokens in zip(rows, documents, strict=True):
        frequencies = Counter(tokens)
        score = 0.0
        for token in query_tokens:
            frequency = frequencies[token]
            if not frequency:
                continue
            idf = math.log(1 + (len(documents) - document_frequency[token] + 0.5) /
                           (document_frequency[token] + 0.5))
            denominator = frequency + 1.5 * (1 - 0.75 + 0.75 * len(tokens) / max(1, average_length))
            score += idf * frequency * 2.5 / denominator
        if score > 0:
            scored.append((score, chunk, title, resource_type))
    scored.sort(key=lambda item: item[0], reverse=True)
    return [_chunk_result(chunk, title, resource_type, round(score, 4), "bm25")
            for score, chunk, title, resource_type in scored[:limit]]


def _chunk_result(
    chunk, title: str, resource_type: str | None, score: float, source: str,
) -> dict:
    return {
        "chunk_id": chunk.id,
        "resource_id": chunk.resource_id,
        "title": title,
        "resource_type": resource_type,
        "position": chunk.position,
        "text": chunk.text,
        "score": score,
        "retrieval_source": source,
        "block_type": chunk.block_type,
        "heading_path": chunk.heading_path,
        "page_number": chunk.page_number,
        "slide_number": chunk.slide_number,
    }


def _complete_sentence_tail(text: str, limit: int = 360) -> str:
    parts = [part.strip() for part in re.split(r"(?<=[。！？；])", text) if part.strip()]
    complete = [part for part in parts[:-1] if re.search(r"[。！？；]$", part)]
    value = "".join(complete[-2:])
    return value[-limit:].lstrip() if value else ""


def _complete_sentence_head(text: str, limit: int = 360) -> str:
    parts = [part.strip() for part in re.split(r"(?<=[。！？；])", text) if part.strip()]
    complete = [part for part in parts if re.search(r"[。！？；]$", part)]
    value = "".join(complete[:2])
    return value[:limit].rstrip() if value else ""


def _attach_context_windows(db: Session, results: list[dict]) -> list[dict]:
    """Attach readable neighbour context without changing the ranked chunk ID."""
    from app.models import ResourceChunk

    resource_ids = {int(item["resource_id"]) for item in results if item.get("resource_id")}
    if not resource_ids:
        return results
    rows = db.scalars(
        select(ResourceChunk)
        .where(ResourceChunk.resource_id.in_(resource_ids))
        .order_by(ResourceChunk.resource_id, ResourceChunk.position)
    ).all()
    grouped: dict[int, list[Any]] = {}
    locations: dict[int, tuple[list[Any], int]] = {}
    for chunk in rows:
        grouped.setdefault(chunk.resource_id, []).append(chunk)
    for values in grouped.values():
        for index, chunk in enumerate(values):
            locations[chunk.id] = (values, index)

    contextualized: list[dict] = []
    for result in results:
        value = dict(result)
        current_text = str(value.get("text") or "").strip()
        pieces: list[str] = []
        location = locations.get(value.get("chunk_id"))
        if location:
            siblings, index = location
            current = siblings[index]
            if index > 0 and siblings[index - 1].heading_path == current.heading_path:
                previous = _complete_sentence_tail(siblings[index - 1].text)
                if previous and not current_text.startswith(previous):
                    pieces.append(previous)
            pieces.append(current_text)
            if index + 1 < len(siblings) and siblings[index + 1].heading_path == current.heading_path:
                following = _complete_sentence_head(siblings[index + 1].text)
                if following and not current_text.endswith(following):
                    pieces.append(following)
        else:
            pieces.append(current_text)
        heading = str(value.get("heading_path") or "").strip()
        context = "\n\n".join(piece for piece in pieces if piece)
        value["context_text"] = f"{heading}\n\n{context}" if heading else context
        contextualized.append(value)
    return contextualized


def _dense_search(offering_id: int, query: str, limit: int) -> list[dict]:
    from qdrant_client import models

    settings = get_settings()
    try:
        query_vector = embed_texts([_chapter_aliases(query)])[0]
        result = _qdrant_client(timeout=5).query_points(
            collection_name=settings.qdrant_collection,
            query=query_vector,
            query_filter=models.Filter(must=[
                models.FieldCondition(key="offering_id", match=models.MatchValue(value=offering_id))
            ]),
            limit=limit,
        ).points
    except Exception:
        return []
    return [{
        "chunk_id": item.payload.get("chunk_id"),
        "resource_id": item.payload["resource_id"],
        "title": item.payload["title"],
        "resource_type": item.payload.get("resource_type"),
        "position": item.payload["position"],
        "text": item.payload["text"],
        "score": round(float(item.score), 4),
        "retrieval_source": "dense",
        "block_type": item.payload.get("block_type", "paragraph"),
        "heading_path": item.payload.get("heading_path"),
        "page_number": item.payload.get("page_number"),
        "slide_number": item.payload.get("slide_number"),
    } for item in result]


def _rrf(dense: Sequence[dict], lexical: Sequence[dict]) -> list[dict]:
    merged: dict[str, dict] = {}
    for source in (dense, lexical):
        for rank, item in enumerate(source):
            key = str(item.get("chunk_id") or f'{item["resource_id"]}:{item["position"]}')
            if key not in merged:
                merged[key] = {**item, "rrf_score": 0.0, "retrieval_sources": []}
            merged[key]["rrf_score"] += 1 / (60 + rank + 1)
            merged[key]["retrieval_sources"].append(item["retrieval_source"])
            if item["retrieval_source"] == "dense":
                merged[key]["score"] = item["score"]
    return sorted(merged.values(), key=lambda item: item["rrf_score"], reverse=True)


@lru_cache(maxsize=1)
def _load_reranker():
    if not get_settings().enable_local_reranker:
        return None
    try:
        from huggingface_hub import snapshot_download
        from sentence_transformers import CrossEncoder

        model_name = get_settings().reranker_model
        try:
            # Prefer an already downloaded model. Hugging Face otherwise performs
            # remote HEAD requests even when all artifacts exist in the cache,
            # which makes an offline evaluation wait through several retries.
            local_snapshot = snapshot_download(model_name, local_files_only=True)
            return CrossEncoder(
                local_snapshot,
                max_length=512,
                local_files_only=True,
            )
        except Exception:
            # Production may legitimately need the first download. Evaluation
            # still uses ``required=True`` and fails if this attempt cannot load.
            return CrossEncoder(model_name, max_length=512)
    except Exception:
        return None


def _rerank(query: str, candidates: list[dict]) -> list[dict]:
    results, _ = _rerank_with_diagnostics(query, candidates)
    return results


def _rerank_with_diagnostics(
    query: str, candidates: list[dict], *, required: bool = False,
) -> tuple[list[dict], dict[str, Any]]:
    model = _load_reranker()
    diagnostics: dict[str, Any] = {
        "reranker_requested": True,
        "reranker_applied": False,
        "reranker_model": get_settings().reranker_model,
        "reranker_device": None,
        "reranked_candidates": 0,
    }
    if model is None:
        if required:
            raise RuntimeError(
                "Reranker 未加载；请执行 uv sync --extra reranker 并确认模型可用"
            )
        return candidates, diagnostics
    diagnostics["reranker_device"] = str(
        getattr(model, "device", getattr(getattr(model, "model", None), "device", "unknown"))
    )
    if not candidates:
        diagnostics["reranker_applied"] = True
        return candidates, diagnostics
    try:
        scores = model.predict(
            [(_chapter_aliases(query), _candidate_retrieval_text(item)) for item in candidates[:12]],
            show_progress_bar=False,
        )
        for item, score in zip(candidates[:12], scores, strict=True):
            item["rerank_score"] = float(score)
        head = sorted(candidates[:12], key=lambda item: item["rerank_score"], reverse=True)
        diagnostics.update({
            "reranker_applied": True,
            "reranked_candidates": len(head),
        })
        return head + candidates[12:], diagnostics
    except Exception as exc:
        if required:
            raise RuntimeError(f"Reranker 执行失败：{exc}") from exc
        return candidates, diagnostics


def retrieve_course(
    offering_id: int, query: str, *, mode: RetrievalMode = "hybrid_rerank",
    limit: int = 5, db: Session | None = None, require_reranker: bool = False,
) -> RetrievalResult:
    """Retrieve course chunks through an explicit, measurable ablation path."""
    from app.db import SessionLocal

    if mode not in {"bm25", "dense", "hybrid", "hybrid_rerank"}:
        raise ValueError(f"不支持的检索模式：{mode}")
    if limit < 1:
        raise ValueError("limit 必须大于 0")
    if require_reranker and mode != "hybrid_rerank":
        raise ValueError("require_reranker 只能用于 hybrid_rerank 模式")

    owns_session = db is None
    session = db or SessionLocal()
    started = perf_counter()
    timings = {"dense": 0.0, "bm25": 0.0, "fusion": 0.0, "rerank": 0.0}
    diagnostics: dict[str, Any] = {
        "dense_candidates": 0,
        "bm25_candidates": 0,
        "fused_candidates": 0,
        "reranker_requested": mode == "hybrid_rerank",
        "reranker_applied": False,
        "reranker_model": get_settings().reranker_model,
        "reranker_device": None,
        "reranked_candidates": 0,
    }
    try:
        dense: list[dict] = []
        lexical: list[dict] = []
        if mode in {"dense", "hybrid", "hybrid_rerank"}:
            phase = perf_counter()
            dense = _dense_search(offering_id, query, 20)
            timings["dense"] = round((perf_counter() - phase) * 1000, 3)
            diagnostics["dense_candidates"] = len(dense)
        if mode in {"bm25", "hybrid", "hybrid_rerank"}:
            phase = perf_counter()
            lexical = _bm25_search(session, offering_id, query, 20)
            timings["bm25"] = round((perf_counter() - phase) * 1000, 3)
            diagnostics["bm25_candidates"] = len(lexical)

        if mode == "bm25":
            results = _apply_source_preferences(query, lexical)
        elif mode == "dense":
            results = _apply_source_preferences(query, dense)
        else:
            phase = perf_counter()
            results = _apply_source_preferences(query, _rrf(dense, lexical))
            timings["fusion"] = round((perf_counter() - phase) * 1000, 3)
            diagnostics["fused_candidates"] = len(results)
            if mode == "hybrid_rerank":
                phase = perf_counter()
                results, rerank_diagnostics = _rerank_with_diagnostics(
                    query, results, required=require_reranker,
                )
                results = _apply_source_preferences(query, results)
                timings["rerank"] = round((perf_counter() - phase) * 1000, 3)
                diagnostics.update(rerank_diagnostics)
        timings["total"] = round((perf_counter() - started) * 1000, 3)
        return RetrievalResult(
            mode=mode, results=_attach_context_windows(session, results[:limit]), timings_ms=timings,
            diagnostics=diagnostics,
        )
    finally:
        if owns_session:
            session.close()


def retrieve_course_variants(
    offering_id: int, query: str, *, limit: int = 10, db: Session | None = None,
    require_reranker: bool = False,
) -> dict[RetrievalMode, RetrievalResult]:
    """Run shared retrieval stages once and expose all four ablation variants."""
    from app.db import SessionLocal

    owns_session = db is None
    session = db or SessionLocal()
    try:
        phase = perf_counter()
        dense = _dense_search(offering_id, query, 20)
        dense_ms = round((perf_counter() - phase) * 1000, 3)
        phase = perf_counter()
        lexical = _bm25_search(session, offering_id, query, 20)
        bm25_ms = round((perf_counter() - phase) * 1000, 3)
        phase = perf_counter()
        dense = _apply_source_preferences(query, dense)
        lexical = _apply_source_preferences(query, lexical)
        fused = _apply_source_preferences(query, _rrf(dense, lexical))
        fusion_ms = round((perf_counter() - phase) * 1000, 3)
        phase = perf_counter()
        reranked, rerank_diagnostics = _rerank_with_diagnostics(
            query, [dict(item) for item in fused], required=require_reranker,
        )
        reranked = _apply_source_preferences(query, reranked)
        rerank_ms = round((perf_counter() - phase) * 1000, 3)
        base_diagnostics = {
            "dense_candidates": len(dense), "bm25_candidates": len(lexical),
            "fused_candidates": len(fused), "reranker_requested": False,
            "reranker_applied": False, "reranker_model": get_settings().reranker_model,
            "reranker_device": None, "reranked_candidates": 0,
        }
        return {
            "bm25": RetrievalResult(
                mode="bm25", results=lexical[:limit],
                timings_ms={"dense": 0.0, "bm25": bm25_ms, "fusion": 0.0,
                            "rerank": 0.0, "total": bm25_ms},
                diagnostics=dict(base_diagnostics),
            ),
            "dense": RetrievalResult(
                mode="dense", results=dense[:limit],
                timings_ms={"dense": dense_ms, "bm25": 0.0, "fusion": 0.0,
                            "rerank": 0.0, "total": dense_ms},
                diagnostics=dict(base_diagnostics),
            ),
            "hybrid": RetrievalResult(
                mode="hybrid", results=fused[:limit],
                timings_ms={"dense": dense_ms, "bm25": bm25_ms, "fusion": fusion_ms,
                            "rerank": 0.0, "total": round(dense_ms + bm25_ms + fusion_ms, 3)},
                diagnostics=dict(base_diagnostics),
            ),
            "hybrid_rerank": RetrievalResult(
                mode="hybrid_rerank", results=reranked[:limit],
                timings_ms={"dense": dense_ms, "bm25": bm25_ms, "fusion": fusion_ms,
                            "rerank": rerank_ms,
                            "total": round(dense_ms + bm25_ms + fusion_ms + rerank_ms, 3)},
                diagnostics={**base_diagnostics, **rerank_diagnostics},
            ),
        }
    finally:
        if owns_session:
            session.close()


def search_course(
    offering_id: int, query: str, limit: int = 5, db: Session | None = None,
) -> list[dict]:
    """Course-scoped dense + BM25 retrieval, fused with RRF and optional reranking."""
    return retrieve_course(
        offering_id, query, mode="hybrid_rerank", limit=limit, db=db,
    ).results


def search_course_supporting(
    offering_id: int, query: str, limit: int = 4, db: Session | None = None,
    resource_ids: Sequence[int] | None = None,
) -> list[dict]:
    """Retrieve a small, chapter-specific set from textbook/courseware only.

    Course-map generation deliberately uses lexical retrieval here. The resources
    have already been embedded during indexing, while this local lookup avoids one
    extra embedding request per chapter and keeps route generation bounded.
    """
    from app.db import SessionLocal
    from app.models import CourseResource, ProcessingStatus, ResourceChunk

    owns_session = db is None
    session = db or SessionLocal()
    try:
        statement = (
            select(ResourceChunk, CourseResource)
            .join(CourseResource, CourseResource.id == ResourceChunk.resource_id)
            .where(
                ResourceChunk.offering_id == offering_id,
                CourseResource.deleted_at.is_(None),
                CourseResource.processing_status == ProcessingStatus.ready,
                CourseResource.resource_type.in_(["textbook", "courseware"]),
            )
            .order_by(ResourceChunk.id)
        )
        if resource_ids:
            statement = statement.where(CourseResource.id.in_(list(resource_ids)))
        rows = session.execute(statement).all()
        query_tokens = _tokenize(query)
        if not rows or not query_tokens:
            return []

        documents = [_tokenize(chunk.text) for chunk, _ in rows]
        document_frequency: Counter[str] = Counter()
        for tokens in documents:
            document_frequency.update(set(tokens))
        average_length = sum(len(tokens) for tokens in documents) / max(1, len(documents))
        scored: list[tuple[float, Any, Any]] = []
        for (chunk, resource), tokens in zip(rows, documents, strict=True):
            frequencies = Counter(tokens)
            score = 0.0
            for token in query_tokens:
                frequency = frequencies[token]
                if not frequency:
                    continue
                idf = math.log(
                    1 + (len(documents) - document_frequency[token] + 0.5) /
                    (document_frequency[token] + 0.5)
                )
                denominator = frequency + 1.5 * (
                    1 - 0.75 + 0.75 * len(tokens) / max(1, average_length)
                )
                score += idf * frequency * 2.5 / denominator
            if score > 0:
                scored.append((score, chunk, resource))
        scored.sort(key=lambda item: item[0], reverse=True)
        return [{
            **_chunk_result(
                chunk, resource.title, resource.resource_type,
                round(score, 4), "bm25_support",
            ),
            "resource_type": resource.resource_type,
        } for score, chunk, resource in scored[:limit]]
    finally:
        if owns_session:
            session.close()


def delete_resource_vectors(resource_id: int) -> None:
    from qdrant_client import models

    settings = get_settings()
    try:
        _qdrant_client(timeout=5).delete(
            collection_name=settings.qdrant_collection,
            points_selector=models.FilterSelector(filter=models.Filter(must=[
                models.FieldCondition(key="resource_id", match=models.MatchValue(value=resource_id))
            ])),
            wait=True,
        )
    except Exception:
        return
