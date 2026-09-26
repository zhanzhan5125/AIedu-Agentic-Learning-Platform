from __future__ import annotations

from io import BytesIO

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.integrations.rag import (
    DocumentBlock, _apply_source_preferences, _chapter_aliases, _clean_pdf_line,
    _complete_sentence_head, _complete_sentence_tail, _estimated_tokens,
    _heading_info, _section_index_blocks,
    _structured_retrieval_text, chunk_blocks, extract_blocks, index_resource,
)


def test_structured_chunks_keep_locator_and_token_limit():
    blocks = [DocumentBlock(
        block_type="paragraph", heading_path="第一章 > 长文本", page_number=3,
        text=("递归用于把复杂问题分解为规模更小的同类问题。" * 120),
    )]
    values = chunk_blocks(blocks)
    assert len(values) > 1
    assert all(_estimated_tokens(item["text"]) <= 800 for item in values)
    assert all(item["page_number"] == 3 for item in values)
    assert all(item["heading_path"] == "第一章 > 长文本" for item in values)


def test_docx_and_pptx_extract_structured_blocks():
    document = Document()
    document.add_heading("数据结构", level=1)
    document.add_paragraph("线性表是有限序列。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "结构"
    table.cell(0, 1).text = "特点"
    docx_bytes = BytesIO()
    document.save(docx_bytes)
    docx_blocks, _ = extract_blocks(
        docx_bytes.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert any(item.block_type == "heading" and item.heading_path == "数据结构" for item in docx_blocks)
    assert any(item.block_type == "table" and "特点" in item.text for item in docx_blocks)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "排序算法"
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    box.text = "快速排序通过枢轴完成分区。"
    pptx_bytes = BytesIO()
    presentation.save(pptx_bytes)
    pptx_blocks, slide_count = extract_blocks(
        pptx_bytes.getvalue(),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    assert slide_count == 1
    assert any(item.slide_number == 1 and item.heading_path == "排序算法" for item in pptx_blocks)
    assert any(item.block_type == "slide" and "快速排序" in item.text for item in pptx_blocks)


def test_plain_style_chinese_headings_are_recovered_from_docx():
    document = Document()
    document.add_paragraph("第三章 C程序的流程设计")
    document.add_paragraph("3.1 结构化算法的性质与结构")
    document.add_paragraph("顺序、选择和循环是三种基本结构。")
    buffer = BytesIO()
    document.save(buffer)

    blocks, _ = extract_blocks(
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert [item.heading_path for item in blocks if item.block_type == "heading"] == [
        "第三章 C程序的流程设计",
        "第三章 C程序的流程设计 > 3.1 结构化算法的性质与结构",
    ]
    body = next(item for item in blocks if item.block_type == "paragraph")
    assert body.heading_path.endswith("3.1 结构化算法的性质与结构")


def test_pdf_noise_heading_and_page_boundaries_are_handled():
    assert _clean_pdf_line("www.linuxidc.com") == ""
    assert _clean_pdf_line("Linux公社(LinuxIDC.com) 是技术网站") == ""
    assert _heading_info("第3章 控制流") == (1, "第3章 控制流")
    assert _heading_info("3.2 if-else 语句") == (2, "3.2 if-else 语句")

    values = chunk_blocks([
        DocumentBlock("paragraph", "第一页内容。", "第一章", page_number=1),
        DocumentBlock("paragraph", "第二页内容。", "第一章", page_number=2),
    ])
    assert len(values) == 2
    assert [item["page_number"] for item in values] == [1, 2]

    overlapped = chunk_blocks([
        DocumentBlock("paragraph", "上一页完整语义。" * 20, "第一章", page_number=1),
        DocumentBlock("paragraph", "下一页继续说明。" * 20, "第一章", page_number=2),
    ])
    assert "上一页完整语义" in overlapped[1]["text"]
    assert "下一页继续说明" in overlapped[1]["text"]


def test_context_window_helpers_only_use_complete_sentences():
    assert _complete_sentence_tail("残缺开头。第一句完整。第二句完整。未完") == "第一句完整。第二句完整。"
    assert _complete_sentence_head("第一句完整。第二句完整。未完") == "第一句完整。第二句完整。"


def test_structured_retrieval_text_and_source_preferences():
    value = _structured_retrieval_text(
        title="C语言教材", resource_type="textbook",
        heading_path="第3章 控制流", text="if-else 用于条件分支。",
    )
    assert "资料类型：教材" in value
    assert "章节：第3章 第三章 控制流" in value
    assert _chapter_aliases("教材第三章讲什么") == "教材第三章 第3章讲什么"

    candidates = [
        {"chunk_id": 1, "resource_type": "syllabus"},
        {"chunk_id": 2, "resource_type": "textbook"},
        {"chunk_id": 3, "resource_type": "syllabus"},
    ]
    assert [item["chunk_id"] for item in _apply_source_preferences("教材第三章", candidates)] == [2]
    assert [item["resource_type"] for item in _apply_source_preferences(
        "结合教材和大纲说明", candidates,
    )[:2]] == ["textbook", "syllabus"]


def test_project_specific_course_requirements_infer_syllabus_and_chapter():
    candidates = [
        {"chunk_id": 1, "resource_type": "textbook", "heading_path": "第2章 数据"},
        {"chunk_id": 2, "resource_type": "syllabus", "heading_path": "第二章 数据"},
        {"chunk_id": 3, "resource_type": "textbook", "heading_path": "第1章 导言"},
        {"chunk_id": 4, "resource_type": "syllabus", "heading_path": "第一章 程序设计初步"},
    ]

    values = _apply_source_preferences(
        "课程第一章要求什么，教材如何说明？", candidates,
    )

    assert [item["chunk_id"] for item in values[:2]] == [3, 4]


def test_section_index_is_built_from_real_child_headings():
    values = _section_index_blocks([
        DocumentBlock("heading", "第5章 指针与数组", "第5章 指针与数组", page_number=96),
        DocumentBlock("heading", "5.1 指针与地址", "第5章 指针与数组 > 5.1 指针与地址", page_number=96),
        DocumentBlock("heading", "5.2 指针与函数参数", "第5章 指针与数组 > 5.2 指针与函数参数", page_number=98),
    ])
    assert len(values) == 1
    assert values[0].block_type == "section_index"
    assert values[0].heading_path == "第5章 指针与数组"
    assert "5.1 指针与地址" in values[0].text
    assert "5.2 指针与函数参数" in values[0].text


def test_index_resource_batches_qdrant_writes_before_stale_cleanup(monkeypatch):
    events = []

    class FakeClient:
        def scroll(self, **kwargs):
            return [type("Record", (), {"id": 999})()], None

        def upsert(self, **kwargs):
            events.append(("upsert", len(kwargs["points"])))

        def delete(self, **kwargs):
            events.append(("delete", len(kwargs["points_selector"].points)))

    monkeypatch.setattr("app.integrations.rag.embed_texts", lambda texts: [[0.1] for _ in texts])
    monkeypatch.setattr("app.integrations.rag._qdrant_client", lambda: FakeClient())
    monkeypatch.setattr("app.integrations.rag._ensure_collection", lambda client: None)

    count = index_resource(
        resource_id=9,
        offering_id=2,
        title="批量教材",
        resource_type="textbook",
        indexed_chunks=[{"id": index + 1, "text": f"内容{index}"} for index in range(130)],
    )

    assert count == 130
    assert events[:3] == [("upsert", 64), ("upsert", 64), ("upsert", 2)]
    assert events[-1] == ("delete", 1)
